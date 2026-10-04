from docxtpl import DocxTemplate
from docx import Document
import csv


with open('books.csv', newline='', encoding='utf-8') as books:
    books_reader = csv.reader(books, delimiter=';')
    next(books_reader)
    lines = sorted(
        [line for line in books_reader if int(line[-1]) > 1],
        key=lambda x: int(x[0])
    )

    doc = Document('book_card.docx')
    fragment_text = '\n'.join(p.text for p in doc.paragraphs)

    names = []
    doc_template = DocxTemplate('book_card.docx')

    for number, line in enumerate(lines):
        context = {
            "fragment_content": fragment_text,
            "ID": line[0],
            "Title": line[1],
            "Author": line[2],
            "Year": line[3],
            "Copies": line[4],
        }
        doc_template.render(context)
        file_name = f'library{number}.docx'
        doc_template.save(file_name)
        names.append(file_name)

    result_doc = Document()
    
    for name in names:
        src_doc = Document(name)
        for par in src_doc.paragraphs:
            result_doc._element.body.append(par._element)
    
    result_doc.save('library.docx')